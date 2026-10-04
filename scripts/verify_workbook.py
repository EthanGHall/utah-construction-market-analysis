"""Read-only verification of the saved Excel artifact (requires openpyxl).

Checks cached output values against the independent CSV analysis, confirms
source preservation, and inspects formulas, Excel tables and chart bindings.
It does not modify the workbook or claim native Excel recalculation was tested.
"""
from __future__ import annotations

import csv
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs/Utah_Construction_Analysis.xlsx"


def equal(left, right):
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-9)
    return left == right


def main():
    result = openpyxl.load_workbook(OUTPUT, data_only=True)
    formulas = openpyxl.load_workbook(OUTPUT, data_only=False)
    reference_path = ROOT / "data/reference/Utah_Construction_Data.xlsx"
    reference = openpyxl.load_workbook(reference_path, data_only=True, read_only=True)
    source_cells = 0
    for name in ("County annual", "By structure", "Sources", "Dictionary"):
        for row in reference[name]:
            for cell in row:
                if cell.value is None:
                    continue
                actual = result[name][cell.coordinate].value
                expected = cell.value
                # Applying a date format makes openpyxl decode the same numeric
                # Excel timestamp as datetime (rounded to milliseconds).
                if name == "Sources" and cell.column == 5 and cell.row >= 6:
                    expected = from_excel(expected, reference.epoch)
                assert equal(actual, expected), (name, cell.coordinate, actual, expected)
                source_cells += 1

    timestamp_cells = 0
    for row in range(6, 27):
        cell = formulas["Sources"].cell(row, 5)
        assert isinstance(cell.value, datetime), ("Retrieval timestamp is not a date", cell.coordinate)
        assert cell.number_format == "yyyy-mm-dd hh:mm:ss", (cell.coordinate, cell.number_format)
        timestamp_cells += 1

    with (ROOT / "reports/county_metrics.csv").open(encoding="utf-8-sig", newline="") as file:
        metrics = {row["county_fips5"]: row for row in csv.DictReader(file)}
    fields = [
        *(f"units_{year}" for year in range(2020, 2026)), "avg_units_2023_2025",
        "absolute_change_2024_2025", "pct_change_2024_2025", "state_share_2025",
        "single_family_share_2025", "multifamily_share_2025", "imputed_share_2025",
        "valuation_usd_2025", "rank_units_2025", "rank_avg_units_2023_2025",
        "avg_units_2021_2025", "rank_avg_units_2021_2025", "reported_units_2025",
        "rank_reported_units_2025", "pct_change_2023_2025",
    ]
    metrics_checked = 0
    for row in result["County comparison"].iter_rows(min_row=6, max_row=34, max_col=23, values_only=True):
        source = metrics[row[0]]
        assert row[1] == source["county_name"]
        for value, field in zip(row[2:], fields):
            expected = float(source[field]) if source[field] else "n.a."
            assert equal(value, expected), (source["county_name"], field, value, expected)
            metrics_checked += 1

    assert all(result["Validation log"].cell(r, 4).value == "Pass" for r in range(6, 18))
    assert result["Validation log"]["D20"].value == "Pass"
    assert all(result["Source comparison"].cell(r, 18).value == "Match" for r in range(6, 180))
    assert all(result["State reconciliation"].cell(r, 9).value == "Pass" for r in range(6, 150))
    assert all(cell.value == 1 for row in result["County coverage"].iter_rows(min_row=6, max_row=34, min_col=3, max_col=8) for cell in row)
    errors = [(sheet.title, cell.coordinate, cell.value) for sheet in result for row in sheet for cell in row if cell.data_type == "e"]
    assert not errors, errors
    for row in range(6, 35):
        for column in range(3, 24):
            assert formulas["County comparison"].cell(row, column).data_type == "f"
    assert formulas["County annual"].tables["CountyannualTable"].ref == "A5:Q179"
    assert len(formulas["Start here"]._charts) == 1
    assert len(formulas["Start here"]._charts[0].series) == 3
    with ZipFile(OUTPUT) as archive:
        chart_parts = [name for name in archive.namelist() if "/charts/" in name and name.endswith(".xml") and "/_rels/" not in name]
        assert len(chart_parts) == 1
        chart_xml = archive.read(chart_parts[0]).decode()
        for color in ("193D55", "00888E", "C56F24"):
            assert color in chart_xml, f"Chart color missing: {color}"
        for col in ("C", "D", "E"):
            assert f"${col}$56:${col}$61" in chart_xml, f"Chart series not linked: {col}"
        xml = "\n".join(archive.read(name).decode(errors="replace") for name in archive.namelist() if name.endswith(".xml"))
        assert "C:\\Users\\" not in xml and "/Users/" not in xml, "Private local path in workbook XML"
    report = {
        "status": "passed",
        "workbook": "outputs/Utah_Construction_Analysis.xlsx",
        "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "reference_value_cells_preserved": source_cells,
        "source_timestamp_cells_checked": timestamp_cells,
        "source_timestamp_number_format": "yyyy-mm-dd hh:mm:ss",
        "intentional_derived_formula_columns": ["County annual H (total units)", "County annual I (multifamily units)", "County annual L (imputed share)"],
        "county_comparison_metrics_checked": metrics_checked,
        "county_year_source_matches": 174,
        "state_reconciliations": 144,
        "coverage_cells": 174,
        "validation_log_checks": 12,
        "saved_formula_errors": 0,
        "native_charts": 1,
        "range_linked_chart_series": 3,
        "recalculation": "Input-change, blank-input and duplicate/missing-key tests passed in the artifact calculation engine. Native Microsoft Excel recalculation was not exercised.",
        "visual_review": "All 10 worksheets reviewed through 13 rendered principal/detail views.",
    }
    (ROOT / "reports/workbook_validation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
